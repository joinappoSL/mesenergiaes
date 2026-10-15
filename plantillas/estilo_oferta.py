"""Estilo y piezas comunes de las plantillas Excel de ofertas de MESENERGIA (+Energía).

Cada operadora tiene su propio generador (oferta-avatel/, oferta-o2/) con su tarifario y su
hoja DATOS; aquí vive lo que comparten: la paleta de la marca, los ayudantes de celdas, la
cabecera de DATOS, la compactación de servicios de CALC y la hoja OFERTA que ve el cliente.
"""

from pathlib import Path

from openpyxl.drawing.image import Image
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties

LOGO = Path(__file__).resolve().parent.parent / "public" / "logo.png"

# Paleta de la marca (src/index.css)
NAVY = "0F1E43"
NAVY_700 = "1B3263"
NAVY_100 = "D9E1F0"
NAVY_50 = "EEF2F9"
VERDE = "6CA83C"
VERDE_50 = "F2F8EC"
SOL = "F5C030"
INPUT = "FFF4CC"
AUTO = "F4F6FA"
GRIS = "5B6475"
TEXTO = "2B3445"
BLANCO = "FFFFFF"
FUENTE = "Arial"

EURO = '#,##0.00 "€"'
FECHA = "dd/mm/yyyy"


def font(size=10, bold=False, color=NAVY, italic=False):
    return Font(name=FUENTE, size=size, bold=bold, color=color, italic=italic)


def fill(color):
    return PatternFill("solid", start_color=color, end_color=color)


def side(color=NAVY_100, style="thin"):
    return Side(style=style, color=color)


def borde(color=NAVY_100):
    s = side(color)
    return Border(left=s, right=s, top=s, bottom=s)


def escribir(ws, ref, valor=None, *, f=None, bg=None, al=None, fmt=None, merge=None, bd=None):
    """Escribe una celda (y opcionalmente la combina) aplicando estilo a todo el rango."""
    celda = ws[ref]
    if valor is not None:
        celda.value = valor
    rango = f"{ref}:{merge or ref}"
    if merge:
        ws.merge_cells(rango)
    for fila in ws[rango]:
        for c in fila:
            if f:
                c.font = f
            if bg:
                c.fill = fill(bg)
            if al:
                c.alignment = al
            if fmt:
                c.number_format = fmt
            if bd:
                c.border = bd
    return celda


def alto(ws, fila, puntos):
    ws.row_dimensions[fila].height = puntos


def logo(ws, ancla, alto_px):
    img = Image(str(LOGO))
    proporcion = img.width / img.height
    img.height = alto_px
    img.width = round(alto_px * proporcion)
    ws.add_image(img, ancla)


IZQ = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
DER = Alignment(horizontal="right", vertical="center")
CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
DESBLOQUEADA = Protection(locked=False)


def imprimir_a4(ws, area=None, *, apaisado=False, una_pagina=True):
    if area:
        ws.print_area = area
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if apaisado else "portrait"
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1 if una_pagina else 0


def validacion(ws, rango, **kw):
    dv = DataValidation(allow_blank=True, showErrorMessage=True, **kw)
    ws.add_data_validation(dv)
    for parte in rango.split():
        dv.add(parte)
    return dv


def fmt_tel(ref):
    """Fórmula que muestra un teléfono de 9 dígitos como «612 345 678»."""
    return f'IF(LEN({ref})=9,LEFT({ref},3)&" "&MID({ref},4,3)&" "&MID({ref},7,3),{ref}&"")'


# ---------------------------------------------------------------------------
# TARIFAS
# ---------------------------------------------------------------------------
def titulo_tarifas(ws, titulo, fuente, hasta):
    ws.sheet_properties.tabColor = NAVY
    ws.sheet_view.showGridLines = False
    escribir(ws, "A1", titulo, f=font(14, True), merge=f"{hasta}1")
    escribir(ws, "A2", fuente, f=font(9, color=GRIS, italic=True),
             al=Alignment(wrap_text=True, vertical="top"), merge=f"{hasta}2")
    alto(ws, 2, 26)


def cabecera_tabla(ws, fila, textos, desde="A"):
    col0 = ord(desde)
    for i, texto in enumerate(textos):
        escribir(ws, f"{chr(col0 + i)}{fila}", texto, f=font(9, True, BLANCO), bg=NAVY_700,
                 al=CENTRO, bd=borde(NAVY_700))


def titulo_seccion(ws, fila, texto, hasta):
    escribir(ws, f"A{fila}", texto, f=font(11, True, VERDE), merge=f"{hasta}{fila}")
    for c in ws[f"A{fila}:{hasta}{fila}"][0]:
        c.border = Border(bottom=side(VERDE, "medium"))


def filas_texto(ws, fila0, filas, *, alto_fila=None, hasta="C"):
    """Tabla «concepto | valor editable» (datos de empresa y textos de la oferta), valor en B:hasta."""
    for i, (concepto, valor) in enumerate(filas):
        r = fila0 + i
        if alto_fila:
            alto(ws, r, alto_fila)
        escribir(ws, f"A{r}", concepto, f=font(10, True), al=IZQ, bd=borde())
        escribir(ws, f"B{r}", valor, f=font(10), bg=INPUT, al=IZQ, merge=f"{hasta}{r}", bd=borde())


def datos_empresa(operadora):
    return [
        ("Empresa", "MESENERGIA"),
        ("Web", "www.mesenergia.es"),
        ("Email de contacto", None),
        ("Teléfono de contacto", None),
        ("Operadora", operadora),
        ("Texto de presentación",
         f"En MESENERGIA somos distribuidores oficiales de {operadora} y comercializamos ofertas exclusivas "
         "diseñadas específicamente para nuestros clientes, con condiciones preferentes y promociones "
         "especiales no disponibles en otros canales."),
    ]


# ---------------------------------------------------------------------------
# DATOS
# ---------------------------------------------------------------------------
def cabecera_datos(ws, operadora, anchos):
    ws.sheet_properties.tabColor = SOL
    ws.sheet_view.showGridLines = False
    for col, ancho in zip("ABCDEFG", anchos):
        ws.column_dimensions[col].width = ancho

    for r in (1, 2, 3):
        alto(ws, r, 20)
    logo(ws, "B1", 72)
    escribir(ws, "C1", "MESENERGIA · +Energía", f=font(9, True, VERDE), merge="G1",
             al=Alignment(vertical="bottom"))
    escribir(ws, "C2", f"Generador de ofertas {operadora}", f=font(17, True), merge="G2",
             al=Alignment(vertical="center"))
    escribir(ws, "C3", "Rellena los datos durante la llamada · la hoja OFERTA se actualiza sola · expórtala a PDF",
             f=font(9, color=GRIS), merge="G3", al=Alignment(vertical="top"))
    alto(ws, 4, 4)
    for c in ws["A4:G4"][0]:
        c.fill = fill(VERDE)

    alto(ws, 6, 46)
    escribir(
        ws,
        "B6",
        "CÓMO SE USA:  1) Rellena las celdas amarillas (las grises se calculan solas; usa Tab para saltar de una a otra).  "
        "2) Revisa la hoja OFERTA.  3) Exporta a PDF solo la hoja OFERTA → Excel: Archivo › Exportar › Crear PDF "
        "(o Imprimir › «Imprimir hojas activas» › Guardar como PDF). Google Sheets: Archivo › Descargar › PDF › Hoja actual.",
        f=font(9, color=NAVY_700),
        bg=NAVY_50,
        al=Alignment(wrap_text=True, vertical="center", indent=1),
        merge="G6",
    )
    escribir(ws, "B7", "Celda a rellenar", f=font(8, color=GRIS), bg=INPUT, al=CENTRO, bd=borde())
    escribir(ws, "C7", "Se calcula sola", f=font(8, color=GRIS), bg=AUTO, al=CENTRO, bd=borde())
    escribir(ws, "D7", "Ejemplo relleno: sobrescríbelo con los datos del cliente.",
             f=font(8, color=GRIS, italic=True), merge="G7", al=Alignment(vertical="center"))


def seccion(ws, fila, texto, hasta="G"):
    escribir(ws, f"B{fila}", texto, f=font(10, True, BLANCO), bg=NAVY,
             al=Alignment(vertical="center", indent=1), merge=f"{hasta}{fila}")
    alto(ws, fila, 22)


def etiqueta(ws, fila, texto, bold=False):
    escribir(ws, f"B{fila}", texto, f=font(10, bold), al=IZQ, bd=borde())


def entrada(ws, ref, valor=None, *, merge=None, fmt=None, al=IZQ):
    celda = escribir(ws, ref, valor, f=font(10, True, NAVY_700), bg=INPUT, al=al, fmt=fmt,
                     merge=merge, bd=borde(SOL))
    for fila in ws[f"{ref}:{merge or ref}"]:
        for c in fila:
            c.protection = DESBLOQUEADA
    return celda


def automatica(ws, ref, valor, *, merge=None, fmt=None, al=IZQ, f=None):
    return escribir(ws, ref, valor, f=f or font(10), bg=AUTO, al=al, fmt=fmt, merge=merge, bd=borde())


def nota(ws, ref, texto, merge="G"):
    fila = ref[1:]
    escribir(ws, ref, texto, f=font(8, color=GRIS, italic=True),
             al=Alignment(vertical="center", wrap_text=True, indent=1), merge=f"{merge}{fila}")


def cabecera_lineas(ws, fila, titulo_bono="Bono de datos (solo adicionales)"):
    for col, texto in zip("BCDEFG", ["Línea", "Número (9 dígitos)", titulo_bono,
                                     "Datos", "€/mes extra", "Comprobación"]):
        escribir(ws, f"{col}{fila}", texto, f=font(9, True, BLANCO), bg=NAVY_700, al=CENTRO, bd=borde(NAVY_700))
    alto(ws, fila, 26)


def comprobacion_linea(r, incluida, precio_bono):
    """Fórmula de la columna «Comprobación» de una línea móvil de DATOS (número en C, bono en D)."""
    return (
        f'=IF(AND(C{r}="",D{r}=""),"",'
        f'IF(C{r}="","⚠ Falta el número",'
        f'IF(LEN(C{r})<>9,"⚠ El número debe tener 9 dígitos",'
        f'IF($C$18="","⚠ Elige primero la tarifa",'
        f'IF({incluida},"✔ Incluida en la tarifa",'
        f'IF(D{r}="","⚠ Elige el bono de datos",'
        f'IF(ISNUMBER(IFERROR({precio_bono},"")),"✔ Línea adicional","⚠ Bono no disponible con esta tarifa")))))))'
    )


def estado(ws, fila, formula, rango_comprobacion):
    seccion(ws, fila - 1, "ESTADO")
    escribir(ws, f"B{fila}", formula, f=font(11, True), al=Alignment(vertical="center", indent=1),
             merge=f"G{fila}", bd=borde())
    alto(ws, fila, 28)
    ws.conditional_formatting.add(
        f"B{fila}:G{fila}",
        FormulaRule(formula=[f'LEFT($B${fila},1)="✔"'], fill=fill(VERDE_50), font=Font(color="2F6B1A")),
    )
    ws.conditional_formatting.add(
        f"B{fila}:G{fila}",
        FormulaRule(formula=[f'LEFT($B${fila},1)<>"✔"'], fill=fill("FEF1C7"), font=Font(color="A95F0F")),
    )
    primera = rango_comprobacion.split(":")[0]
    ws.conditional_formatting.add(
        rango_comprobacion,
        FormulaRule(formula=[f'LEFT(${primera[0]}{primera[1:]},1)="⚠"'], font=Font(color="C0392B", bold=True)),
    )


def proteger_datos(ws):
    ws.protection.sheet = True
    ws.protection.formatColumns = False
    ws.protection.formatRows = False


# ---------------------------------------------------------------------------
# CALC (oculta)
# ---------------------------------------------------------------------------
def preparar_calc(c):
    c.sheet_state = "hidden"
    for col, ancho in zip("ABCDE", (22, 16, 26, 90, 8)):
        c.column_dimensions[col].width = ancho
    c["A1"] = "Hoja auxiliar: no editar. La usa la hoja OFERTA."
    c["A1"].font = font(9, True, GRIS)


def calc_lineas(c, filas_datos):
    """Filas 3-7: número, datos y viñeta de cada línea móvil de DATOS; B8 = nº de líneas."""
    c["A2"], c["B2"], c["C2"], c["D2"] = "Línea", "Número", "Datos", "Viñeta"
    for i, r_d in enumerate(filas_datos, start=1):
        r = 2 + i
        c[f"A{r}"] = "Línea principal" if i == 1 else "Línea adicional"
        c[f"B{r}"] = f'=DATOS!$C${r_d}&""'
        c[f"C{r}"] = f'=DATOS!$E${r_d}&""'
        c[f"D{r}"] = f'=IF(B{r}="","",CHAR(10)&"•  "&A{r}&"  "&{fmt_tel(f"B{r}")}&"   ·   "&C{r})'
    c["A8"], c["B8"] = "Nº líneas", '=SUMPRODUCT(--(LEN(B3:B7)>0))'


def calc_servicios(c, servicios):
    """Filas 11-16: servicios candidatos (clave, condición, título, descripción).

    Filas 19-24: los mismos, compactados para que la oferta no tenga huecos.
    """
    c["A10"], c["B10"], c["C10"], c["D10"], c["E10"] = "Servicio", "Incluir", "Título", "Descripción", "Acum."
    for i, (clave, cond, titulo, desc) in enumerate(servicios):
        r = 11 + i
        c[f"A{r}"] = clave
        c[f"B{r}"] = f"=IF({cond},1,0)"
        c[f"C{r}"] = titulo
        c[f"D{r}"] = f"={desc}"
        c[f"E{r}"] = f"=B{r}" if i == 0 else f"=E{r - 1}+B{r}"

    c["A18"], c["B18"], c["C18"], c["D18"] = "Orden", "", "Título", "Descripción"
    for k in range(1, 7):
        r = 18 + k
        c[f"A{r}"] = k
        c[f"C{r}"] = f'=IFERROR(INDEX($C$11:$C$16,MATCH({k},$E$11:$E$16,0)),"")'
        c[f"D{r}"] = f'=IFERROR(INDEX($D$11:$D$16,MATCH({k},$E$11:$E$16,0)),"")'


# ---------------------------------------------------------------------------
# OFERTA: documento para el cliente
# ---------------------------------------------------------------------------
def hoja_oferta(o, *, operadora, prefijo_ref, tarifa, d, t, altos_servicios):
    """Construye la hoja OFERTA.

    tarifa: fórmula (sin «=») con el nombre comercial de la tarifa elegida.
    d: referencias a DATOS (nombre, fecha, asesor, validez, precio).
    t: referencias a TARIFAS (empresa, web, email, tel, operadora, intro, pie).
    La hoja CALC aporta la línea principal (B3), los servicios (C19:D24) y el resumen (D27).
    """
    o.sheet_properties.tabColor = VERDE
    o.sheet_view.showGridLines = False
    o.sheet_view.zoomScale = 100
    for col, ancho in zip("ABCDEFG", (1.3, 25, 1.5, 17, 17, 17, 17)):
        o.column_dimensions[col].width = ancho

    # Cabecera: logo + datos de empresa
    for r in range(1, 5):
        alto(o, r, 17)
    logo(o, "A1", 78)
    escribir(o, "D1", f'={t["empresa"]}&""', f=font(12, True), al=DER, merge="G1")
    escribir(o, "D2", f'={t["web"]}&""', f=font(9, color=GRIS), al=DER, merge="G2")
    escribir(o, "D3", f'={t["email"]}&""', f=font(9, color=GRIS), al=DER, merge="G3")
    escribir(o, "D4", f'={t["tel"]}&""', f=font(9, color=GRIS), al=DER, merge="G4")
    alto(o, 5, 8)
    alto(o, 6, 3)
    for cc in o["A6:G6"][0]:
        cc.fill = fill(VERDE)
    alto(o, 7, 12)

    # Título
    alto(o, 8, 34)
    escribir(o, "A8", "CONDICIONES Y TARIFA PERSONALIZADA", f=font(15, True, BLANCO), bg=NAVY,
             al=Alignment(horizontal="center", vertical="center"), merge="G8")
    alto(o, 9, 14)

    # Cliente / fecha / referencia
    etiq = font(7.5, True, GRIS)
    alto(o, 10, 13)
    escribir(o, "A10", "CLIENTE", f=etiq, al=Alignment(vertical="bottom"), merge="D10")
    escribir(o, "E10", "FECHA", f=etiq, al=Alignment(horizontal="center", vertical="bottom"))
    escribir(o, "F10", "REFERENCIA", f=etiq, al=Alignment(horizontal="center", vertical="bottom"), merge="G10")
    alto(o, 11, 26)
    escribir(o, "A11", f'=IF({d["nombre"]}="","—",{d["nombre"]})', f=font(16, True),
             al=Alignment(vertical="center"), merge="D11")
    fecha = d["fecha"]
    escribir(o, "E11", f'=IF({fecha}="","",{fecha})', f=font(10, True), al=CENTRO, fmt=FECHA)
    escribir(
        o, "F11",
        f'=IF({fecha}="","","{prefijo_ref}-"&YEAR({fecha})&RIGHT("0"&MONTH({fecha}),2)&RIGHT("0"&DAY({fecha}),2)'
        f'&IF(CALC!$B$3="","","-"&RIGHT(CALC!$B$3,4)))',
        f=font(10, True), al=CENTRO, merge="G11",
    )
    for cc in o["A11:G11"][0]:
        cc.border = Border(bottom=side(NAVY_100))
    alto(o, 12, 18)
    asesor, validez = d["asesor"], d["validez"]
    escribir(o, "A12", f'=IF({asesor}="","","Asesor comercial: "&{asesor})', f=font(9, color=GRIS),
             al=Alignment(vertical="center"), merge="D12")
    escribir(o, "E12", f'=IF({validez}="","","Válida hasta")', f=font(9, color=GRIS),
             al=Alignment(horizontal="right", vertical="center"))
    escribir(o, "F12", f'=IF({validez}="","",{validez})', f=font(9, True, GRIS), fmt=FECHA,
             al=Alignment(horizontal="center", vertical="center"), merge="G12")
    alto(o, 13, 10)

    # Presentación
    alto(o, 14, 54)
    escribir(
        o, "A14",
        f'={t["intro"]}&" Pásate a "&{t["operadora"]}&" y disfruta de la siguiente oferta, preparada en exclusiva para ti."',
        f=font(10, color=TEXTO), al=Alignment(wrap_text=True, vertical="center", horizontal="justify"),
        merge="G14",
    )
    alto(o, 15, 12)

    # Bloque de precio
    for r, h in ((16, 8), (17, 14), (18, 30), (19, 16), (20, 8)):
        alto(o, r, h)
        for cc in o[f"A{r}:G{r}"][0]:
            cc.fill = fill(NAVY)
    for r in range(16, 21):
        o[f"A{r}"].fill = fill(VERDE)
    escribir(o, "B17", f"TU TARIFA {operadora}", f=font(8, True, SOL), al=Alignment(vertical="bottom", indent=1),
             merge="D17")
    escribir(o, "B18", f'=IFERROR({tarifa},"Selecciona una tarifa")',
             f=font(16, True, BLANCO), al=Alignment(vertical="center", indent=1), merge="D18")
    escribir(o, "B19", "=CALC!$D$27", f=font(8.5, color=NAVY_100), al=Alignment(vertical="top", indent=1),
             merge="E19")
    escribir(o, "E17", "CUOTA MENSUAL", f=font(8, True, SOL),
             al=Alignment(horizontal="right", vertical="bottom", indent=1), merge="G17")
    escribir(o, "E18", f'={d["precio"]}', f=font(24, True, SOL),
             al=Alignment(horizontal="right", vertical="center", indent=1), fmt=EURO, merge="G18")
    escribir(o, "F19", "al mes · IVA incluido", f=font(8.5, color=BLANCO),
             al=Alignment(horizontal="right", vertical="top", indent=1), merge="G19")
    alto(o, 21, 16)

    # Servicios incluidos
    alto(o, 22, 22)
    escribir(o, "A22", "SERVICIOS INCLUIDOS EN TU OFERTA", f=font(11, True), al=Alignment(vertical="center"),
             merge="G22")
    for cc in o["A22:G22"][0]:
        cc.border = Border(bottom=side(VERDE, "medium"))
    alto(o, 23, 6)

    for k, h in enumerate(altos_servicios, start=1):
        r = 23 + k
        alto(o, r, h)
        escribir(o, f"B{r}", f"=CALC!$C${18 + k}", f=font(10.5, True),
                 al=Alignment(vertical="center", wrap_text=True, indent=1))
        escribir(o, f"D{r}", f"=CALC!$D${18 + k}", f=font(9.5, color=TEXTO),
                 al=Alignment(vertical="center", wrap_text=True), merge=f"G{r}")
    o.conditional_formatting.add("A24:A29", FormulaRule(formula=['$B24<>""'], fill=fill(VERDE)))
    o.conditional_formatting.add("B24:B29", FormulaRule(formula=['$B24<>""'], fill=fill(VERDE_50)))
    o.conditional_formatting.add("A24:G29", FormulaRule(formula=['$B24<>""'], border=Border(bottom=side(NAVY_100))))

    # Cierre
    alto(o, 30, 14)
    alto(o, 31, 40)
    email, tel = t["email"], t["tel"]
    escribir(
        o, "A31",
        f'=IF({asesor}="","Nuestro equipo",{asesor})&" te acompañará durante todo el proceso de alta, portabilidad '
        f'e instalación para que no tengas que preocuparte de nada."'
        f'&IF(AND({email}="",{tel}=""),""," Para cualquier consulta, contacta con nosotros en "'
        f'&{email}&IF(AND({email}<>"",{tel}<>"")," o en el ","")&{tel}&".")',
        f=font(10, color=TEXTO), bg=NAVY_50,
        al=Alignment(wrap_text=True, vertical="center", horizontal="left", indent=1), merge="G31",
    )
    alto(o, 32, 10)
    alto(o, 33, 26)
    escribir(o, "A33", f'={t["pie"]}&""', f=font(7.5, color=GRIS, italic=True),
             al=Alignment(wrap_text=True, vertical="top"), merge="G33")

    imprimir_a4(o, "A1:G33")
    o.print_options.horizontalCentered = True
    o.page_margins.left = o.page_margins.right = 0.55
    o.page_margins.top = 0.5
    o.page_margins.bottom = 0.6
    o.page_margins.footer = 0.3
    o.oddFooter.center.text = "MESENERGIA  ·  +Energía  ·  Nos importa tu ahorro"
    o.oddFooter.center.size = 8
    o.oddFooter.center.font = "Arial"
    o.oddFooter.center.color = GRIS
    o.protection.sheet = True


def guardar(wb, salida):
    wb.active = 0
    wb.calculation.fullCalcOnLoad = True
    wb.save(salida)
    print(f"Plantilla generada: {salida}")
