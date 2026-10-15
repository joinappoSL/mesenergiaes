"""Genera la plantilla Excel de ofertas AVATEL de MESENERGIA (+Energía).

Hojas:
  DATOS   -> se rellena durante la llamada (celdas amarillas).
  OFERTA  -> documento corporativo para el cliente; se actualiza solo y se exporta a PDF.
  TARIFAS -> precios y textos del tarifario AVATEL (editable cuando cambie el tarifario).
  CALC    -> cálculos auxiliares (oculta).

Uso:  python3 generar_plantilla.py [salida.xlsx]
"""

import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment

AQUI = Path(__file__).resolve().parent
sys.path.insert(0, str(AQUI.parent))

from estilo_oferta import (  # noqa: E402
    CENTRO, DER, EURO, FECHA, GRIS, INPUT, IZQ, NAVY,
    alto, automatica, borde, cabecera_datos, cabecera_lineas, cabecera_tabla, calc_lineas, calc_servicios,
    comprobacion_linea, datos_empresa, entrada, escribir, estado, etiqueta, filas_texto, fmt_tel, font,
    guardar, hoja_oferta, imprimir_a4, nota, preparar_calc, proteger_datos, seccion, titulo_seccion,
    titulo_tarifas, validacion,
)

SALIDA = Path(sys.argv[1]) if len(sys.argv) > 1 else AQUI / "Oferta_AVATEL_MesEnergia.xlsx"

wb = Workbook()
datos = wb.active
datos.title = "DATOS"
oferta = wb.create_sheet("OFERTA")
tarifas = wb.create_sheet("TARIFAS")
calc = wb.create_sheet("CALC")

# ---------------------------------------------------------------------------
# TARIFAS: tarifario AVATEL + datos de empresa + textos de la oferta
# ---------------------------------------------------------------------------
t = tarifas
for col, ancho in zip("ABCDEF", (30, 28, 74, 14, 26, 16)):
    t.column_dimensions[col].width = ancho
titulo_tarifas(
    t,
    "TARIFAS AVATEL · Oferta comercial cliente nuevo",
    "Fuente: tarifario AVATEL Telecom + Energía (vigencia 18/08/2026 – 15/10/2026). "
    "Las celdas amarillas se pueden actualizar cuando cambie el tarifario: la oferta se recalcula sola.",
    "F",
)

# Paquetes (filas 6-9)
titulo_seccion(t, 4, "PAQUETES", "F")
cabecera_tabla(
    t,
    5,
    [
        "Tarifa",
        "Nombre en la oferta",
        "Qué incluye (referencia interna)",
        "Líneas móviles incluidas",
        "Datos de las líneas incluidas",
        "Precio €/mes (IVA incl.)",
    ],
)
alto(t, 5, 30)
PAQUETES = [
    ("Conecta Fibra", "Conecta Fibra prime",
     "Fibra 1000 Mb · Router WiFi 6 · Instalación gratis · 12 meses de Amazon Prime",
     0, "—", 19.99),
    ("Conecta Sencillo", "Conecta Sencillo prime",
     "Fibra 1000 Mb · Router WiFi 6 · Instalación gratis · 1 línea móvil 100 GB acumulables "
     "con llamadas ilimitadas · 12 meses de Amazon Prime",
     1, "100 GB acumulables", 24.99),
    ("Conecta Doble", "Conecta Doble prime",
     "Fibra 1000 Mb · Router WiFi 6 · Instalación gratis · 2 líneas móviles 5G 100 GB acumulables "
     "con llamadas ilimitadas · 12 meses de Amazon Prime",
     2, "100 GB acumulables (5G)", 34.99),
    ("Conecta Sin Límites", "Conecta Sin Límites prime",
     "Fibra 1000 Mb · Router WiFi 6 · Instalación gratis · 2 líneas móviles 5G con datos ilimitados "
     "y llamadas ilimitadas · 12 meses de Amazon Prime",
     2, "Datos ilimitados (5G)", 39.99),
]
for i, (clave, nombre, incluye, lineas, datos_inc, precio) in enumerate(PAQUETES):
    r = 6 + i
    alto(t, r, 30)
    escribir(t, f"A{r}", clave, f=font(10, True), al=IZQ, bd=borde())
    for col, valor, fmt, al in (
        ("B", nombre, None, IZQ),
        ("C", incluye, None, IZQ),
        ("D", lineas, "0", CENTRO),
        ("E", datos_inc, None, IZQ),
        ("F", precio, EURO, CENTRO),
    ):
        escribir(t, f"{col}{r}", valor, f=font(10), bg=INPUT, al=al, fmt=fmt, bd=borde())

# Bonos líneas adicionales (cabecera fila 12, matriz B13:E16)
titulo_seccion(t, 11, "BONOS PARA LÍNEAS MÓVILES ADICIONALES (€/mes) · todas con llamadas ilimitadas · máx. 5 líneas por cliente", "F")
BONOS = ["25 GB", "50 GB", "100 GB acumulables", "Datos ilimitados"]
cabecera_tabla(t, 12, ["Tarifa"] + BONOS)
alto(t, 12, 22)
MATRIZ = {
    "Conecta Fibra": [5.00, None, None, 12.00],
    "Conecta Sencillo": [5.00, 7.50, None, 12.00],
    "Conecta Doble": [None, None, 5.00, None],
    "Conecta Sin Límites": [None, None, None, 10.00],
}
for i, (clave, *_rest) in enumerate(PAQUETES):
    r = 13 + i
    escribir(t, f"A{r}", clave, f=font(10, True), al=IZQ, bd=borde())
    for j, precio in enumerate(MATRIZ[clave]):
        col = "BCDE"[j]
        escribir(t, f"{col}{r}", precio if precio is not None else "—",
                 f=font(10, color=NAVY if precio is not None else GRIS), bg=INPUT,
                 al=CENTRO, fmt=EURO, bd=borde())
escribir(t, "A17", "«—» = bono no disponible con esa tarifa.", f=font(8, color=GRIS, italic=True), merge="F17")

# Extras (filas 21-24)
titulo_seccion(t, 19, "EXTRAS (€/mes)", "F")
cabecera_tabla(t, 20, ["Concepto", "€/mes", "Nota"])
EXTRAS = [
    ("Línea de teléfono fijo", 5.00, "1.000 min a fijos y 1.000 min a móviles nacionales"),
    ("Televisión CLICtv (con deco)", 5.00, ""),
    ("Canal Local CLICtv", 0.00, "Gratuito hasta el 31/12/2026 según el tarifario; actualiza el precio cuando cambie"),
    ("Conecta Mesh (por dispositivo)", 3.99, ""),
]
for i, (concepto, precio, texto_nota) in enumerate(EXTRAS):
    r = 21 + i
    escribir(t, f"A{r}", concepto, f=font(10, True), al=IZQ, bd=borde())
    escribir(t, f"B{r}", precio, f=font(10), bg=INPUT, al=CENTRO, fmt=EURO, bd=borde())
    escribir(t, f"C{r}", texto_nota, f=font(9, color=GRIS), al=IZQ, bd=borde())

# Datos de empresa (filas 27-32)
titulo_seccion(t, 26, "DATOS DE EMPRESA (aparecen en la oferta)", "F")
filas_texto(t, 27, datos_empresa("AVATEL"))
alto(t, 32, 44)

# Textos de los servicios (filas 35-40)
titulo_seccion(t, 34, "TEXTOS DE LA OFERTA (lo que lee el cliente)", "F")
filas_texto(t, 35, [
    ("Fibra óptica",
     "Conexión a Internet de alta velocidad mediante red de fibra óptica de 1 Gb (1000 Mb), "
     "con router WiFi 6 incluido e instalación gratuita."),
    ("Teléfono fijo (tras el número)",
     "con 1.000 minutos a fijos y 1.000 minutos a móviles nacionales."),
    ("Televisión",
     "Servicio de televisión CLICtv con decodificador incluido, para disfrutar de tus canales "
     "y contenidos favoritos."),
    ("Canal Local",
     "Incluye además el Canal Local de CLICtv, con LALIGA HYPERMOTION y PRIMERA RFEF."),
    ("Amazon Prime",
     "12 meses de Amazon Prime incluidos con tu tarifa, conforme a las condiciones del "
     "proveedor del servicio."),
    ("Nota al pie",
     "Precios con IVA incluido. Servicios prestados por AVATEL, sujetos a cobertura y a la "
     "validación de los datos del titular por parte de la operadora."),
], alto_fila=30)
imprimir_a4(t, apaisado=True, una_pagina=False)

T_PAQ = "TARIFAS!$A$6:$A$9"
T_NOMBRE = "TARIFAS!$B$6:$B$9"
T_INCLUYE = "TARIFAS!$C$6:$C$9"
T_LINEAS = "TARIFAS!$D$6:$D$9"
T_DATOS = "TARIFAS!$E$6:$E$9"
T_PRECIO = "TARIFAS!$F$6:$F$9"
T_BONOS = "TARIFAS!$B$12:$E$12"
T_MATRIZ_PAQ = "TARIFAS!$A$13:$A$16"
T_MATRIZ = "TARIFAS!$B$13:$E$16"
T_FIJO, T_TV, T_CANAL, T_MESH = "TARIFAS!$B$21", "TARIFAS!$B$22", "TARIFAS!$B$23", "TARIFAS!$B$24"
T_EMPRESA, T_WEB, T_EMAIL, T_TEL, T_OPERADORA, T_INTRO = (f"TARIFAS!$B${r}" for r in range(27, 33))
TX_FIBRA, TX_FIJO, TX_TV, TX_CANAL, TX_AMAZON, TX_PIE = (f"TARIFAS!$B${r}" for r in range(35, 41))

# ---------------------------------------------------------------------------
# DATOS: lo que rellena el comercial en la llamada
# ---------------------------------------------------------------------------
d = datos
cabecera_datos(d, "AVATEL", (2.5, 34, 20, 26, 18, 14, 40))
for r in range(9, 45):
    alto(d, r, 20)

seccion(d, 9, "1 · DATOS DE LA OFERTA")
etiqueta(d, 10, "Fecha de la oferta")
entrada(d, "C10", date(2026, 10, 13), fmt=FECHA, al=CENTRO)
nota(d, "D10", "Formato dd/mm/aaaa")
etiqueta(d, 11, "Asesor comercial")
entrada(d, "C11", "Gabriela", merge="D11")
nota(d, "E11", "Nombre que verá el cliente en la oferta")
etiqueta(d, 12, "Oferta válida hasta (opcional)")
entrada(d, "C12", None, fmt=FECHA, al=CENTRO)
nota(d, "D12", "Si lo dejas vacío no aparece en la oferta")

seccion(d, 14, "2 · CLIENTE")
etiqueta(d, 15, "Nombre completo", bold=True)
entrada(d, "C15", "María López Fernández", merge="E15")

seccion(d, 17, "3 · TARIFA AVATEL")
etiqueta(d, 18, "Tarifa", bold=True)
entrada(d, "C18", "Conecta Sin Límites", merge="D18")
escribir(d, "E18", "Precio tarifario", f=font(9, color=GRIS), al=DER)
automatica(d, "F18", f"=IFERROR(INDEX({T_PRECIO},MATCH($C$18,{T_PAQ},0)),\"\")", fmt=EURO, al=CENTRO)
etiqueta(d, 19, "Qué incluye")
automatica(d, "C19", f"=IFERROR(INDEX({T_INCLUYE},MATCH($C$18,{T_PAQ},0)),\"Elige una tarifa\")",
           merge="G19", f=font(9))
alto(d, 19, 30)
etiqueta(d, 20, "Líneas móviles incluidas en la tarifa")
automatica(d, "C20", f"=IFERROR(INDEX({T_LINEAS},MATCH($C$18,{T_PAQ},0)),0)", fmt="0", al=CENTRO)
nota(d, "D20", "Las primeras líneas de la tabla de abajo van incluidas; el resto son adicionales con bono")

seccion(d, 22, "4 · LÍNEAS MÓVILES  (portabilidad · máximo 5 por cliente)")
cabecera_lineas(d, 23)

FILAS_LINEAS = range(24, 29)
EJEMPLO_LINEAS = [("612345678", None), ("623456789", None), ("634567890", "Datos ilimitados"), (None, None), (None, None)]
for i, r in enumerate(FILAS_LINEAS, start=1):
    etiqueta(d, r, "Línea principal" if i == 1 else f"Línea adicional {i - 1}", bold=(i == 1))
    numero, bono = EJEMPLO_LINEAS[i - 1]
    entrada(d, f"C{r}", numero, fmt="@", al=CENTRO)
    entrada(d, f"D{r}", bono, al=CENTRO)
    incluida = f"{i}<=$C$20"
    precio_bono = f"INDEX({T_MATRIZ},MATCH($C$18,{T_MATRIZ_PAQ},0),MATCH(D{r},{T_BONOS},0))"
    automatica(
        d, f"E{r}",
        f'=IF(C{r}="","",IF({incluida},INDEX({T_DATOS},MATCH($C$18,{T_PAQ},0)),D{r}))',
        al=CENTRO, f=font(9),
    )
    automatica(
        d, f"F{r}",
        f'=IF(OR(C{r}="",D{r}="",{incluida}),0,N(IFERROR({precio_bono},0)))',
        fmt=EURO, al=CENTRO,
    )
    automatica(d, f"G{r}", comprobacion_linea(r, incluida, precio_bono), f=font(9))

seccion(d, 30, "5 · TELÉFONO FIJO")
etiqueta(d, 31, "Número fijo a portar")
entrada(d, "C31", "944123456", fmt="@", al=CENTRO)
nota(d, "D31", "Déjalo vacío si el cliente no lleva línea fija")

seccion(d, 33, "6 · EXTRAS")
etiqueta(d, 34, "Televisión CLICtv (con deco)")
entrada(d, "C34", "Sí", al=CENTRO)
nota(d, "D34", "Sí / No")
etiqueta(d, 35, "Canal Local CLICtv")
entrada(d, "C35", "No", al=CENTRO)
nota(d, "D35", "Sí / No · requiere CLICtv")
etiqueta(d, 36, "Conecta Mesh (nº de dispositivos)")
entrada(d, "C36", 1, fmt="0", al=CENTRO)
nota(d, "D36", "0 a 5 · amplía la cobertura WiFi")
etiqueta(d, 37, "Otras ventajas para el cliente (opcional)")
entrada(d, "C37", None, merge="G37", al=Alignment(wrap_text=True, vertical="center", indent=1))
alto(d, 37, 36)

seccion(d, 39, "7 · PRECIO")
etiqueta(d, 40, "Precio orientativo según tarifario")
automatica(
    d, "C40",
    f"=IFERROR(INDEX({T_PRECIO},MATCH($C$18,{T_PAQ},0)),0)+SUM(F24:F28)"
    f'+IF(C31<>"",{T_FIJO},0)+IF(C34="Sí",{T_TV},0)+IF(AND(C34="Sí",C35="Sí"),{T_CANAL},0)'
    f"+N(C36)*{T_MESH}",
    fmt=EURO, al=CENTRO,
)
nota(d, "D40", "Tarifa + líneas adicionales + fijo + extras, según la hoja TARIFAS (solo referencia interna)")
etiqueta(d, 41, "PRECIO DE NUESTRA OFERTA (€/mes, IVA incl.)", bold=True)
entrada(d, "C41", 59.90, fmt=EURO, al=CENTRO)
nota(d, "D41", "Es el precio que verá el cliente. Si lo dejas vacío se usará el orientativo")
etiqueta(d, 42, "Precio que aparece en el PDF")
automatica(d, "C42", "=IF(ISNUMBER(C41),C41,C40)", fmt=EURO, al=CENTRO, f=font(11, True))

estado(
    d, 45,
    '=IF(C15="","⚠ Falta el nombre completo del cliente",'
    'IF(C18="","⚠ Elige la tarifa AVATEL",'
    'IF(SUMPRODUCT(--(LEFT(G24:G28,1)="⚠"))>0,"⚠ Revisa las líneas móviles (columna Comprobación)",'
    'IF(AND(C31<>"",LEN(C31)<>9),"⚠ El teléfono fijo debe tener 9 dígitos",'
    'IF(AND(C35="Sí",C34<>"Sí"),"⚠ El Canal Local requiere CLICtv con deco",'
    'IF(NOT(ISNUMBER(C41)),"ℹ Sin precio propio: el PDF mostrará el precio orientativo",'
    '"✔ Oferta lista: ve a la hoja OFERTA y expórtala a PDF"))))))',
    "G24:G28",
)

validacion(d, "C10 C12", type="date", operator="greaterThan", formula1="43831",
           errorTitle="Fecha", error="Introduce una fecha válida (dd/mm/aaaa).")
validacion(d, "C18", type="list", formula1=T_PAQ,
           errorTitle="Tarifa", error="Elige una tarifa de la lista.")
validacion(d, "D24:D28", type="list", formula1=T_BONOS,
           errorTitle="Bono", error="Elige un bono de la lista.")
validacion(d, "C24:C28 C31", type="textLength", operator="equal", formula1="9",
           errorTitle="Teléfono", error="El número debe tener 9 dígitos, sin espacios ni prefijo.")
validacion(d, "C34:C35", type="list", formula1='"Sí,No"',
           errorTitle="Sí / No", error="Elige Sí o No.")
validacion(d, "C36", type="whole", operator="between", formula1="0", formula2="5",
           errorTitle="Mesh", error="Entre 0 y 5 dispositivos.")
validacion(d, "C41", type="decimal", operator="greaterThanOrEqual", formula1="0",
           errorTitle="Precio", error="Introduce el precio mensual en euros (ej. 49,90).")

imprimir_a4(d, "A1:G45")
proteger_datos(d)

D_FECHA, D_ASESOR, D_VALIDEZ, D_NOMBRE, D_TARIFA = "DATOS!$C$10", "DATOS!$C$11", "DATOS!$C$12", "DATOS!$C$15", "DATOS!$C$18"
D_FIJO, D_TV, D_CANAL, D_MESH, D_OTROS, D_PRECIO = (
    "DATOS!$C$31", "DATOS!$C$34", "DATOS!$C$35", "DATOS!$C$36", "DATOS!$C$37", "DATOS!$C$42")

# ---------------------------------------------------------------------------
# CALC: textos de la oferta y compactación de servicios (hoja oculta)
# ---------------------------------------------------------------------------
c = calc
preparar_calc(c)
calc_lineas(c, FILAS_LINEAS)
calc_servicios(c, [
    ("Fibra", "TRUE",
     "Fibra óptica 1 Gb",
     f'{TX_FIBRA}&IF(N({D_MESH})>0," Incluye "&{D_MESH}&IF({D_MESH}=1," dispositivo Conecta Mesh"," dispositivos Conecta Mesh")'
     f'&" para ampliar la cobertura WiFi en toda la vivienda.","")'),
    ("Móvil", "$B$8>0",
     "Telefonía móvil",
     'IF($B$8=1,"Línea móvil en modalidad de portabilidad, con llamadas ilimitadas a destinos nacionales:",'
     '"Líneas móviles en modalidad de portabilidad, todas con llamadas ilimitadas a destinos nacionales:")'
     '&D3&D4&D5&D6&D7'),
    ("Fijo", f'{D_FIJO}<>""',
     "Telefonía fija",
     f'"Línea de teléfono fijo en modalidad de portabilidad con el número "&{fmt_tel(D_FIJO)}&", "&{TX_FIJO}'),
    ("TV", f'{D_TV}="Sí"',
     "Televisión",
     f'{TX_TV}&IF({D_CANAL}="Sí"," "&{TX_CANAL},"")'),
    ("Amazon", "TRUE",
     "Amazon Prime",
     f"{TX_AMAZON}&\"\""),
    ("Ventajas", f'{D_OTROS}<>""',
     "Ventajas +Energía",
     f'{D_OTROS}&""'),
])

# Resumen de la tarifa (fila 27)
c["A27"] = "Resumen"
c["D27"] = (
    '="Fibra 1 Gb"&IF($B$8>0,"  ·  "&$B$8&IF($B$8=1," línea móvil"," líneas móviles"),"")'
    f'&IF({D_FIJO}<>"","  ·  Teléfono fijo","")&IF({D_TV}="Sí","  ·  Televisión","")'
    '&"  ·  Amazon Prime"'
)
c.protection.sheet = True

# ---------------------------------------------------------------------------
# OFERTA: documento para el cliente
# ---------------------------------------------------------------------------
hoja_oferta(
    oferta,
    operadora="AVATEL",
    prefijo_ref="AV",
    tarifa=f"INDEX({T_NOMBRE},MATCH({D_TARIFA},{T_PAQ},0))",
    d={"nombre": D_NOMBRE, "fecha": D_FECHA, "asesor": D_ASESOR, "validez": D_VALIDEZ, "precio": D_PRECIO},
    t={"empresa": T_EMPRESA, "web": T_WEB, "email": T_EMAIL, "tel": T_TEL, "operadora": T_OPERADORA,
       "intro": T_INTRO, "pie": TX_PIE},
    altos_servicios=(50, 92, 50, 48, 44, 44),
)

guardar(wb, SALIDA)
