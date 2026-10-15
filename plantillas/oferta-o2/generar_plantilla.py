"""Genera la plantilla Excel de ofertas O2 de MESENERGIA (+Energía).

Hojas:
  DATOS   -> se rellena durante la llamada (celdas amarillas).
  OFERTA  -> documento corporativo para el cliente; se actualiza solo y se exporta a PDF.
  TARIFAS -> precios y textos del tarifario O2 (editable cuando cambie el tarifario).
  CALC    -> cálculos auxiliares (oculta).

Uso:  python3 generar_plantilla.py [salida.xlsx]
"""

import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment

AQUI = Path(__file__).resolve().parent
sys.path.insert(0, str(AQUI.parent))

from estilo_oferta import (  # noqa: E402
    CENTRO, DER, EURO, FECHA, GRIS, INPUT, IZQ, NAVY,
    alto, automatica, borde, cabecera_datos, cabecera_lineas, cabecera_tabla, calc_lineas, calc_servicios,
    comprobacion_linea, datos_empresa, entrada, escribir, estado, etiqueta, filas_texto, fmt_tel, font,
    guardar, hoja_oferta, imprimir_a4, nota, preparar_calc, proteger_datos, seccion, titulo_seccion,
    titulo_tarifas, validacion,
)

SALIDA = Path(sys.argv[1]) if len(sys.argv) > 1 else AQUI / "Oferta_O2_MesEnergia.xlsx"

wb = Workbook()
datos = wb.active
datos.title = "DATOS"
oferta = wb.create_sheet("OFERTA")
tarifas = wb.create_sheet("TARIFAS")
calc = wb.create_sheet("CALC")

# ---------------------------------------------------------------------------
# TARIFAS: tarifario O2 + datos de empresa + textos de la oferta
# ---------------------------------------------------------------------------
t = tarifas
for col, ancho in zip("ABCDEFGHIJKL", (56, 30, 18, 10, 11, 12, 12, 12, 10, 10, 13, 16)):
    t.column_dimensions[col].width = ancho
titulo_tarifas(
    t,
    "TARIFAS O2",
    "Fuente: tarifario O2 (tarifas vigentes desde el 8 de julio de 2026). Precios con IVA incluido. "
    "Las celdas amarillas se pueden actualizar cuando cambie el tarifario: la oferta se recalcula sola.",
    "L",
)

# Paquetes (filas 6-23)
titulo_seccion(t, 4, "PAQUETES", "L")
cabecera_tabla(t, 5, [
    "Tarifa (lista de la hoja DATOS)",
    "Nombre en la oferta",
    "Tipo",
    "Fibra",
    "Líneas móviles incluidas",
    "Datos línea 1",
    "Datos línea 2",
    "Movistar Plus+",
    "Netflix",
    "Disney+",
    "Precio €/mes (IVA incl.)",
    "Bonos de líneas adicionales",
])
alto(t, 5, 30)
FMV, SOLO_MOVIL = "Fibra y móvil", "Solo móvil"
PAQUETES = [
    # clave, nombre en la oferta, tipo, fibra, líneas, datos 1, datos 2, M+, Netflix, Disney+, precio, bonos
    ("Fibra 600 Mb + 35 GB + Movistar Plus+", "Fibra 600 Mb + Móvil 35 GB", "Fibra, móvil y TV",
     "600 Mb", 1, "35 GB", None, "Sí", "No", "No", 38, FMV),
    ("Fibra 1 Gb + 350 GB + Movistar Plus+", "Fibra 1 Gb + Móvil 350 GB", "Fibra, móvil y TV",
     "1 Gb", 1, "350 GB", None, "Sí", "No", "No", 50, FMV),
    ("Fibra 600 Mb + 60 GB + Movistar Plus+ y Disney+", "Fibra 600 Mb + Móvil 60 GB", "Fibra, móvil y TV",
     "600 Mb", 1, "60 GB", None, "Sí", "No", "Sí", 45, FMV),
    ("Fibra 1 Gb + 375 GB + Movistar Plus+ y Disney+", "Fibra 1 Gb + Móvil 375 GB", "Fibra, móvil y TV",
     "1 Gb", 1, "375 GB", None, "Sí", "No", "Sí", 56, FMV),
    ("Fibra 600 Mb + 60 GB + Movistar Plus+ y Netflix", "Fibra 600 Mb + Móvil 60 GB", "Fibra, móvil y TV",
     "600 Mb", 1, "60 GB", None, "Sí", "Sí", "No", 47, FMV),
    ("Fibra 1 Gb + 375 GB + Movistar Plus+ y Netflix", "Fibra 1 Gb + Móvil 375 GB", "Fibra, móvil y TV",
     "1 Gb", 1, "375 GB", None, "Sí", "Sí", "No", 58, FMV),
    ("Fibra 600 Mb + 60 GB + Movistar Plus+, Netflix y Disney+", "Fibra 600 Mb + Móvil 60 GB", "Fibra, móvil y TV",
     "600 Mb", 1, "60 GB", None, "Sí", "Sí", "Sí", 52, FMV),
    ("Fibra 1 Gb + 375 GB + Movistar Plus+, Netflix y Disney+", "Fibra 1 Gb + Móvil 375 GB", "Fibra, móvil y TV",
     "1 Gb", 1, "375 GB", None, "Sí", "Sí", "Sí", 62, FMV),
    ("Fibra 600 Mb + 60 GB", "Fibra 600 Mb + Móvil 60 GB", "Fibra y móvil",
     "600 Mb", 1, "60 GB", None, "No", "No", "No", 35, FMV),
    ("Fibra 600 Mb + 10 GB + 40 GB (2 líneas)", "Fibra 600 Mb + 10 GB + 40 GB", "Fibra y móvil",
     "600 Mb", 2, "10 GB", "40 GB", "No", "No", "No", 35, FMV),
    ("Fibra 1 Gb + 120 GB", "Fibra 1 Gb + Móvil 120 GB", "Fibra y móvil",
     "1 Gb", 1, "120 GB", None, "No", "No", "No", 38, FMV),
    ("Solo Fibra 300 Mb", "Fibra 300 Mb", "Solo fibra", "300 Mb", 0, None, None, "No", "No", "No", 23, SOLO_MOVIL),
    ("Solo Fibra 600 Mb", "Fibra 600 Mb", "Solo fibra", "600 Mb", 0, None, None, "No", "No", "No", 27, SOLO_MOVIL),
    ("Solo Fibra 1 Gb", "Fibra 1 Gb", "Solo fibra", "1 Gb", 0, None, None, "No", "No", "No", 31, SOLO_MOVIL),
    ("Solo Móvil 50 GB", "Móvil 5G+ 50 GB", "Solo móvil", None, 1, "50 GB", None, "No", "No", "No", 7, SOLO_MOVIL),
    ("Solo Móvil 100 GB", "Móvil 5G+ 100 GB", "Solo móvil", None, 1, "100 GB", None, "No", "No", "No", 10, SOLO_MOVIL),
    ("Solo Móvil 150 GB", "Móvil 5G+ 150 GB", "Solo móvil", None, 1, "150 GB", None, "No", "No", "No", 15, SOLO_MOVIL),
    ("Solo Móvil 300 GB", "Móvil 5G+ 300 GB", "Solo móvil", None, 1, "300 GB", None, "No", "No", "No", 20, SOLO_MOVIL),
]
P0, P1 = 6, 6 + len(PAQUETES) - 1
for i, (clave, *resto) in enumerate(PAQUETES):
    r = P0 + i
    alto(t, r, 20)
    escribir(t, f"A{r}", clave, f=font(10, True), al=IZQ, bd=borde())
    for col, valor in zip("BCDEFGHIJKL", resto):
        fmt = EURO if col == "K" else ("0" if col == "E" else None)
        al = IZQ if col in "BCL" else CENTRO
        escribir(t, f"{col}{r}", valor, f=font(10), bg=INPUT, al=al, fmt=fmt, bd=borde())

# Bonos de líneas adicionales (cabecera fila 26, matriz B27:F28)
titulo_seccion(t, 25, "LÍNEAS MÓVILES ADICIONALES 5G+ (€/mes) · todas con llamadas ilimitadas", "L")
BONOS = ["40 GB", "50 GB", "100 GB", "150 GB", "300 GB"]
cabecera_tabla(t, 26, ["Bonos de líneas adicionales"] + BONOS)
alto(t, 26, 22)
MATRIZ = [
    (FMV, [5, None, None, 10, 15]),
    (SOLO_MOVIL, [None, 7, 10, 15, 20]),
]
for i, (grupo, precios) in enumerate(MATRIZ):
    r = 27 + i
    escribir(t, f"A{r}", grupo, f=font(10, True), al=IZQ, bd=borde())
    for col, precio in zip("BCDEF", precios):
        escribir(t, f"{col}{r}", precio if precio is not None else "—",
                 f=font(10, color=NAVY if precio is not None else GRIS), bg=INPUT,
                 al=CENTRO, fmt=EURO, bd=borde())
escribir(
    t, "A29",
    "«Fibra y móvil»: clientes de Fibra y Móvil con o sin TV. «Solo móvil»: cada línea extra es una tarifa "
    "Solo Móvil (también para clientes de Solo Fibra). «—» = no disponible.",
    f=font(8, color=GRIS, italic=True), merge="L29",
)

# Fijo y fibra de segunda residencia (filas 33 y 37-39)
titulo_seccion(t, 31, "EXTRAS (€/mes)", "L")
cabecera_tabla(t, 32, ["Concepto", "€/mes", "Nota"])
escribir(t, "A33", "Línea de teléfono fijo", f=font(10, True), al=IZQ, bd=borde())
escribir(t, "B33", 0, f=font(10), bg=INPUT, al=CENTRO, fmt=EURO, bd=borde())
escribir(t, "C33", "Con las tarifas de fibra. Ajusta el precio si O2 lo cobra aparte.",
         f=font(9, color=GRIS), al=IZQ, merge="H33", bd=borde())

titulo_seccion(t, 35, "FIBRA SEGUNDA RESIDENCIA (€/mes) · para clientes de Fibra y Móvil con o sin TV", "L")
cabecera_tabla(t, 36, ["Velocidad", "€/mes"])
for i, (velocidad, precio) in enumerate((("300 Mb", 15), ("600 Mb", 20), ("1 Gb", 27))):
    r = 37 + i
    escribir(t, f"A{r}", velocidad, f=font(10, True), al=IZQ, bd=borde())
    escribir(t, f"B{r}", precio, f=font(10), bg=INPUT, al=CENTRO, fmt=EURO, bd=borde())

# Datos de empresa (filas 42-47)
titulo_seccion(t, 41, "DATOS DE EMPRESA (aparecen en la oferta)", "L")
filas_texto(t, 42, datos_empresa("O2"), hasta="H")
alto(t, 47, 32)

# Textos de los servicios (filas 50-55)
titulo_seccion(t, 49, "TEXTOS DE LA OFERTA (lo que lee el cliente)", "L")
filas_texto(t, 50, [
    ("Fibra óptica (tras la velocidad)",
     ", con router WiFi incluido para conectar todos tus dispositivos."),
    ("Telefonía móvil (tras «con»)",
     "llamadas ilimitadas a destinos nacionales"),
    ("Teléfono fijo (tras el número)",
     "con llamadas conforme a las condiciones vigentes del operador."),
    ("Movistar Plus+",
     "Televisión con Movistar Plus+ incluido, para disfrutar de series, cine, documentales y "
     "entretenimiento, conforme a las condiciones del servicio."),
    ("Netflix / Disney+ (tras los nombres)",
     ", conforme a las condiciones del proveedor del servicio."),
    ("Nota al pie",
     "Precios con IVA incluido. Servicios prestados por O2, sujetos a cobertura y a la validación de "
     "los datos del titular por parte de la operadora."),
], alto_fila=30, hasta="H")
imprimir_a4(t, apaisado=True, una_pagina=False)

T_PAQ = f"TARIFAS!$A${P0}:$A${P1}"
T_NOMBRE = f"TARIFAS!$B${P0}:$B${P1}"


def t_col(col):
    return f"TARIFAS!${col}${P0}:${col}${P1}"


T_BONOS = "TARIFAS!$B$26:$F$26"
T_GRUPOS = "TARIFAS!$A$27:$A$28"
T_MATRIZ = "TARIFAS!$B$27:$F$28"
T_FIJO = "TARIFAS!$B$33"
T_2RES, T_2RES_PRECIO = "TARIFAS!$A$37:$A$39", "TARIFAS!$B$37:$B$39"
T_EMPRESA, T_WEB, T_EMAIL, T_TEL, T_OPERADORA, T_INTRO = (f"TARIFAS!$B${r}" for r in range(42, 48))
TX_FIBRA, TX_MOVIL, TX_FIJO, TX_MPLUS, TX_PLATAFORMAS, TX_PIE = (f"TARIFAS!$B${r}" for r in range(50, 56))

# Atributos de la tarifa elegida (CALC, filas 30-42)
C_FIBRA, C_LINEAS, C_DATOS1, C_DATOS2 = "CALC!$B$33", "CALC!$B$34", "CALC!$B$35", "CALC!$B$36"
C_MPLUS, C_NETFLIX, C_DISNEY, C_PRECIO, C_GRUPO, C_INCLUYE = (
    "CALC!$B$37", "CALC!$B$38", "CALC!$B$39", "CALC!$B$40", "CALC!$B$41", "CALC!$B$42")

# ---------------------------------------------------------------------------
# DATOS: lo que rellena el comercial en la llamada
# ---------------------------------------------------------------------------
d = datos
cabecera_datos(d, "O2", (2.5, 34, 22, 30, 18, 16, 36))
for r in range(9, 44):
    alto(d, r, 20)

seccion(d, 9, "1 · DATOS DE LA OFERTA")
etiqueta(d, 10, "Fecha de la oferta")
entrada(d, "C10", date(2026, 10, 15), fmt=FECHA, al=CENTRO)
nota(d, "D10", "Formato dd/mm/aaaa")
etiqueta(d, 11, "Asesor comercial")
entrada(d, "C11", "Gabriela", merge="D11")
nota(d, "E11", "Nombre que verá el cliente en la oferta")
etiqueta(d, 12, "Oferta válida hasta (opcional)")
entrada(d, "C12", None, fmt=FECHA, al=CENTRO)
nota(d, "D12", "Si lo dejas vacío no aparece en la oferta")

seccion(d, 14, "2 · CLIENTE")
etiqueta(d, 15, "Nombre completo", bold=True)
entrada(d, "C15", "María López Fernández", merge="E15")

seccion(d, 17, "3 · TARIFA O2")
etiqueta(d, 18, "Tarifa", bold=True)
entrada(d, "C18", "Fibra 1 Gb + 375 GB + Movistar Plus+ y Netflix", merge="E18")
escribir(d, "F18", "Precio tarifario", f=font(9, color=GRIS), al=DER)
automatica(d, "G18", f'=IF($C$18="","",{C_PRECIO})', fmt=EURO, al=CENTRO)
etiqueta(d, 19, "Qué incluye")
automatica(d, "C19", f'=IF($C$18="","Elige una tarifa",{C_INCLUYE})', merge="G19", f=font(9))
etiqueta(d, 20, "Líneas móviles incluidas en la tarifa")
automatica(d, "C20", f"={C_LINEAS}", fmt="0", al=CENTRO)
nota(d, "D20", "Las primeras líneas de la tabla de abajo van incluidas; el resto son adicionales con bono")

seccion(d, 22, "4 · LÍNEAS MÓVILES 5G+  (portabilidad · hasta 5 líneas)")
cabecera_lineas(d, 23)

FILAS_LINEAS = range(24, 29)
EJEMPLO_LINEAS = [("612345678", None), ("623456789", "150 GB"), (None, None), (None, None), (None, None)]
for i, r in enumerate(FILAS_LINEAS, start=1):
    etiqueta(d, r, "Línea principal" if i == 1 else f"Línea adicional {i - 1}", bold=(i == 1))
    numero, bono = EJEMPLO_LINEAS[i - 1]
    entrada(d, f"C{r}", numero, fmt="@", al=CENTRO)
    entrada(d, f"D{r}", bono, al=CENTRO)
    incluida = f"{i}<=$C$20"
    precio_bono = f"INDEX({T_MATRIZ},MATCH({C_GRUPO},{T_GRUPOS},0),MATCH(D{r},{T_BONOS},0))"
    automatica(
        d, f"E{r}",
        f'=IF(C{r}="","",IF({incluida},{C_DATOS1 if i == 1 else C_DATOS2},D{r}))',
        al=CENTRO, f=font(9),
    )
    automatica(
        d, f"F{r}",
        f'=IF(OR(C{r}="",D{r}="",{incluida}),0,N(IFERROR({precio_bono},0)))',
        fmt=EURO, al=CENTRO,
    )
    automatica(d, f"G{r}", comprobacion_linea(r, incluida, precio_bono), f=font(9))

seccion(d, 30, "5 · TELÉFONO FIJO")
etiqueta(d, 31, "Número fijo a portar")
entrada(d, "C31", "944123456", fmt="@", al=CENTRO)
nota(d, "D31", "Déjalo vacío si el cliente no lleva línea fija (solo con tarifas de fibra)")

seccion(d, 33, "6 · EXTRAS")
etiqueta(d, 34, "Fibra segunda residencia")
entrada(d, "C34", None, al=CENTRO)
nota(d, "D34", "300 Mb / 600 Mb / 1 Gb · solo clientes de Fibra y Móvil · vacío si no lleva")
etiqueta(d, 35, "Otras ventajas para el cliente (opcional)")
entrada(d, "C35", None, merge="G35", al=Alignment(wrap_text=True, vertical="center", indent=1))
alto(d, 35, 36)

seccion(d, 37, "7 · PRECIO")
etiqueta(d, 38, "Precio orientativo según tarifario")
automatica(
    d, "C38",
    f"=N({C_PRECIO})+SUM(F24:F28)"
    f'+IF(AND(C31<>"",{C_FIBRA}<>""),{T_FIJO},0)'
    f"+IFERROR(INDEX({T_2RES_PRECIO},MATCH(C34,{T_2RES},0)),0)",
    fmt=EURO, al=CENTRO,
)
nota(d, "D38", "Tarifa + líneas adicionales + fijo + segunda residencia, según la hoja TARIFAS (solo referencia interna)")
etiqueta(d, 39, "PRECIO DE NUESTRA OFERTA (€/mes, IVA incl.)", bold=True)
entrada(d, "C39", 65.00, fmt=EURO, al=CENTRO)
nota(d, "D39", "Es el precio que verá el cliente. Si lo dejas vacío se usará el orientativo")
etiqueta(d, 40, "Precio que aparece en el PDF")
automatica(d, "C40", "=IF(ISNUMBER(C39),C39,C38)", fmt=EURO, al=CENTRO, f=font(11, True))

estado(
    d, 43,
    '=IF(C15="","⚠ Falta el nombre completo del cliente",'
    'IF(C18="","⚠ Elige la tarifa O2",'
    'IF(AND(C20>0,C24=""),"⚠ Falta el número de la línea principal (va incluida en la tarifa)",'
    'IF(SUMPRODUCT(--(LEFT(G24:G28,1)="⚠"))>0,"⚠ Revisa las líneas móviles (columna Comprobación)",'
    'IF(AND(C31<>"",LEN(C31)<>9),"⚠ El teléfono fijo debe tener 9 dígitos",'
    f'IF(AND(C31<>"",{C_FIBRA}=""),"⚠ El teléfono fijo solo va con tarifas de fibra",'
    f'IF(AND(C34<>"",{C_GRUPO}<>"{FMV}"),"⚠ La fibra de segunda residencia es solo para clientes de Fibra y Móvil",'
    'IF(NOT(ISNUMBER(C39)),"ℹ Sin precio propio: el PDF mostrará el precio orientativo",'
    '"✔ Oferta lista: ve a la hoja OFERTA y expórtala a PDF"))))))))',
    "G24:G28",
)

validacion(d, "C10 C12", type="date", operator="greaterThan", formula1="43831",
           errorTitle="Fecha", error="Introduce una fecha válida (dd/mm/aaaa).")
validacion(d, "C18", type="list", formula1=T_PAQ,
           errorTitle="Tarifa", error="Elige una tarifa de la lista.")
validacion(d, "D24:D28", type="list", formula1=T_BONOS,
           errorTitle="Bono", error="Elige un bono de la lista.")
validacion(d, "C24:C28 C31", type="textLength", operator="equal", formula1="9",
           errorTitle="Teléfono", error="El número debe tener 9 dígitos, sin espacios ni prefijo.")
validacion(d, "C34", type="list", formula1=T_2RES,
           errorTitle="Segunda residencia", error="Elige 300 Mb, 600 Mb o 1 Gb (o déjalo vacío).")
validacion(d, "C39", type="decimal", operator="greaterThanOrEqual", formula1="0",
           errorTitle="Precio", error="Introduce el precio mensual en euros (ej. 49,90).")

imprimir_a4(d, "A1:G43")
proteger_datos(d)

D_FECHA, D_ASESOR, D_VALIDEZ, D_NOMBRE, D_TARIFA = "DATOS!$C$10", "DATOS!$C$11", "DATOS!$C$12", "DATOS!$C$15", "DATOS!$C$18"
D_FIJO, D_2RES, D_OTROS, D_PRECIO = "DATOS!$C$31", "DATOS!$C$34", "DATOS!$C$35", "DATOS!$C$40"

# ---------------------------------------------------------------------------
# CALC: atributos de la tarifa, textos de la oferta y compactación de servicios (hoja oculta)
# ---------------------------------------------------------------------------
c = calc
preparar_calc(c)
calc_lineas(c, FILAS_LINEAS)

CON_FIJO = f'AND({D_FIJO}<>"",$B$33<>"")'
calc_servicios(c, [
    ("Fibra", '$B$33<>""',
     '="Fibra óptica "&$B$33',
     f'"Conexión a Internet de alta velocidad mediante red de fibra óptica de "&$B$33&{TX_FIBRA}'
     f'&IF({D_2RES}<>""," Incluye además fibra de "&{D_2RES}&" para tu segunda residencia.","")'),
    ("Móvil", "$B$8>0",
     "Telefonía móvil 5G+",
     f'IF($B$8=1,"Línea móvil 5G+ en modalidad de portabilidad, con ",'
     f'"Líneas móviles 5G+ en modalidad de portabilidad, todas con ")&{TX_MOVIL}&":"&D3&D4&D5&D6&D7'),
    ("Fijo", CON_FIJO,
     "Telefonía fija",
     f'"Línea de teléfono fijo en modalidad de portabilidad con el número "&{fmt_tel(D_FIJO)}&", "&{TX_FIJO}'),
    ("TV", '$B$37="Sí"',
     "Movistar Plus+",
     f'{TX_MPLUS}&""'),
    ("Plataformas", 'OR($B$38="Sí",$B$39="Sí")',
     "Plataformas de contenidos",
     '"Acceso incluido a "&IF(AND($B$38="Sí",$B$39="Sí"),"Netflix y Disney+",IF($B$38="Sí","Netflix","Disney+"))'
     f'&{TX_PLATAFORMAS}'),
    ("Ventajas", f'{D_OTROS}<>""',
     "Ventajas +Energía",
     f'{D_OTROS}&""'),
])

# Resumen de la tarifa (fila 27): cada parte empieza por «  ·  » y MID quita el primero
c["A27"] = "Resumen"
c["D27"] = (
    '=MID(IF($B$33<>"","  ·  Fibra "&$B$33,"")'
    '&IF($B$8>0,"  ·  "&$B$8&IF($B$8=1," línea móvil"," líneas móviles"),"")'
    f'&IF({CON_FIJO},"  ·  Teléfono fijo","")'
    '&IF($B$37="Sí","  ·  Movistar Plus+","")&IF($B$38="Sí","  ·  Netflix","")&IF($B$39="Sí","  ·  Disney+","")'
    f'&IF({D_2RES}<>"","  ·  Fibra 2ª residencia",""),6,500)'
)

# Atributos de la tarifa elegida (filas 30-42)
c["A30"], c["B30"] = "Posición de la tarifa", f"=IFERROR(MATCH({D_TARIFA},{T_PAQ},0),0)"
for r, (titulo, col) in enumerate((
    ("Nombre", "B"), ("Tipo", "C"), ("Fibra", "D"), ("Líneas incluidas", "E"), ("Datos línea 1", "F"),
    ("Datos línea 2", "G"), ("Movistar Plus+", "H"), ("Netflix", "I"), ("Disney+", "J"), ("Precio", "K"),
    ("Bonos adicionales", "L"),
), start=31):
    c[f"A{r}"] = titulo
    if col in "EK":
        c[f"B{r}"] = f"=IF($B$30=0,0,N(INDEX({t_col(col)},$B$30)))"
    else:
        c[f"B{r}"] = f'=IF($B$30=0,"",INDEX({t_col(col)},$B$30)&"")'
c["A42"] = "Qué incluye"
c["B42"] = (
    '=MID(IF($B$33<>"","  ·  Fibra "&$B$33,"")'
    '&IF($B$35<>"","  ·  Móvil 5G+ "&$B$35&IF($B$36<>""," + Móvil 5G+ "&$B$36,""),"")'
    '&IF($B$37="Sí","  ·  Movistar Plus+","")&IF($B$38="Sí","  ·  Netflix","")&IF($B$39="Sí","  ·  Disney+",""),6,500)'
)
c.protection.sheet = True

# ---------------------------------------------------------------------------
# OFERTA: documento para el cliente
# ---------------------------------------------------------------------------
hoja_oferta(
    oferta,
    operadora="O2",
    prefijo_ref="O2",
    tarifa=f"INDEX({T_NOMBRE},MATCH({D_TARIFA},{T_PAQ},0))",
    d={"nombre": D_NOMBRE, "fecha": D_FECHA, "asesor": D_ASESOR, "validez": D_VALIDEZ, "precio": D_PRECIO},
    t={"empresa": T_EMPRESA, "web": T_WEB, "email": T_EMAIL, "tel": T_TEL, "operadora": T_OPERADORA,
       "intro": T_INTRO, "pie": TX_PIE},
    altos_servicios=(50, 80, 46, 46, 40, 44),
)

guardar(wb, SALIDA)
